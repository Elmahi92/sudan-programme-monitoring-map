from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Presention' / 'UNDP_Sudan_Progress_Report_Jan-Sep2026_draft.pptx'
OUTPUT = ROOT / 'Presention' / 'UNDP_Sudan_Progress_Report_Jan-Sep2026_UNDP_styled.pptx'

NAVY = RGBColor(0x00, 0x3C, 0x5A)
BLUE = RGBColor(0x00, 0x6E, 0xB5)
CYAN = RGBColor(0x5D, 0xD4, 0xF0)
DARK = RGBColor(0x2F, 0x3B, 0x4A)
MUTED = RGBColor(0x68, 0x75, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, color, size=None, bold=None):
    font = run.font
    font.name = 'Arial'
    font.color.rgb = color
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold


def style_text_frame(shape):
    frame = shape.text_frame
    if not frame.paragraphs:
        return
    top = shape.top / 914400
    height = shape.height / 914400
    is_heading = top < 1.25 or height < 0.55
    for paragraph in frame.paragraphs:
        paragraph.font.name = 'Arial'
        paragraph.font.color.rgb = NAVY if is_heading else DARK
        if is_heading:
            paragraph.font.bold = True
        for run in paragraph.runs:
            set_run_font(run, NAVY if is_heading else DARK, bold=True if is_heading else None)
    if is_heading:
        for paragraph in frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT


def add_accent(slide, slide_width):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, Inches(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()


def style_deck():
    prs = Presentation(str(SOURCE))
    for slide in prs.slides:
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE
        add_accent(slide, prs.slide_width)
        for shape in slide.shapes:
            if not getattr(shape, 'has_text_frame', False):
                continue
            if not shape.text.strip():
                continue
            style_text_frame(shape)
    prs.save(str(OUTPUT))


def update_theme():
    replacements = {
        b'A6025C': b'006EB5', b'92248E': b'003C5A', b'DE95C4': b'5DD4F0',
        b'FE4A00': b'F4A261', b'DA002F': b'D1495B', b'FF907A': b'F5B5A5',
        b'4472C4': b'006EB5', b'ED7D31': b'F4A261', b'A5A5A5': b'687580',
        b'FFC000': b'F2C14E', b'5B9BD5': b'5DD4F0', b'70AD47': b'6AA67F',
    }
    temp = OUTPUT.with_suffix('.tmp.pptx')
    with ZipFile(OUTPUT, 'r') as source, ZipFile(temp, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.startswith('ppt/theme/') or item.filename.startswith('ppt/slides/'):
                for old, new in replacements.items():
                    content = content.replace(old, new)
            target.writestr(item, content)
    temp.replace(OUTPUT)


if __name__ == '__main__':
    style_deck()
    update_theme()
    print(f'created={OUTPUT}')
    print(f'slides={len(Presentation(str(OUTPUT)).slides)}')

