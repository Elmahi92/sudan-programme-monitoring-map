"""Build the record-level master GeoJSON from the September workbook."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "Data" / "Master_data_as_sep26.xlsx"
LEGACY_PATH = ROOT / "Geo_Data" / "M_E_masterdata_by_state.geojson"
BOUNDARY_PATH = ROOT / "Geo_Data" / "sdn_admbnda_adm1_cbs_nic_ssa_download.geojson"
OUTPUT_PATH = ROOT / "Geo_Data" / "M_E_masterdata_by_state.geojson"
REPORT_PATH = ROOT / "Data" / "master_data_validation.json"

STATE_ALIASES = {"gedarif": "gedaref", "kartoum": "khartoum"}


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def normalize_state(value: Any) -> str:
    parts = [STATE_ALIASES.get(clean_text(part), clean_text(part)) for part in str(value or "").split("/")]
    return "/".join(part for part in parts if part)


def normalize_number(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    try:
        number = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return value
    return int(number) if number.is_integer() else number


def row_key(properties: dict[str, Any], headers: list[str]) -> tuple[Any, ...]:
    return tuple(properties.get(header, properties.get("Location ") if header == "Location" else None) for header in headers)


def load_workbook_rows() -> tuple[list[str], list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    records = []
    for row in rows[1:]:
        record = {header: normalize_number(value) for header, value in zip(headers, row)}
        records.append(record)
    return headers, records


def load_geojson(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def boundary_lookup(boundaries: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        normalize_state(feature.get("properties", {}).get("adm1_en")): feature.get("geometry")
        for feature in boundaries.get("features", [])
        if feature.get("geometry")
    }


def geometry_for_state(state: Any, boundaries: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    parts = normalize_state(state).split("/")
    geometries = [boundaries.get(part) for part in parts if boundaries.get(part)]
    if len(geometries) == 1:
        return geometries[0]
    if geometries and all(geometry.get("type") == "Polygon" for geometry in geometries):
        return {"type": "MultiPolygon", "coordinates": [geometry["coordinates"] for geometry in geometries]}
    return None


def main() -> int:
    headers, records = load_workbook_rows()
    legacy = load_geojson(LEGACY_PATH)
    boundaries = boundary_lookup(load_geojson(BOUNDARY_PATH))
    legacy_by_key: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for feature in legacy.get("features", []):
        key = row_key(feature.get("properties", {}), headers)
        legacy_by_key.setdefault(key, []).append(feature.get("geometry"))

    output_features = []
    matched_legacy = 0
    fallback_boundary = 0
    unmatched_states = Counter()
    null_geometry_rows = []
    duplicate_keys = [key for key, count in Counter(row_key(record, headers) for record in records).items() if count > 1]

    for row_number, record in enumerate(records, start=2):
        key = row_key(record, headers)
        old_geometries = legacy_by_key.get(key, [])
        geometry = old_geometries.pop(0) if old_geometries else None
        if geometry is not None:
            matched_legacy += 1
        if geometry is None:
            geometry = geometry_for_state(record.get("State"), boundaries)
            if geometry is not None:
                fallback_boundary += 1
        if geometry is None:
            normalized = normalize_state(record.get("State"))
            unmatched_states[normalized or "<NULL>"] += 1
            null_geometry_rows.append(row_number)
        output_features.append({"type": "Feature", "properties": record, "geometry": geometry})

    output = {"type": "FeatureCollection", "features": output_features}
    OUTPUT_PATH.write_text(json.dumps(output, ensure_ascii=True, separators=(",", ":")), encoding="utf-8")
    report = {
        "source_workbook": str(WORKBOOK_PATH.relative_to(ROOT)),
        "source_sheet": "Sheet1",
        "source_records": len(records),
        "output_features": len(output_features),
        "source_headers": headers,
        "legacy_geometry_matches": matched_legacy,
        "boundary_geometry_fallbacks": fallback_boundary,
        "null_geometry_records": len(null_geometry_rows),
        "null_geometry_row_numbers": null_geometry_rows,
        "unmatched_or_null_states": dict(unmatched_states),
        "duplicate_source_records": len(duplicate_keys),
        "note": "Rows are preserved one-for-one; duplicate records are not aggregated.",
    }
    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=True), encoding="utf-8")
    print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())